import os
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side
from docx import Document

def extract_faculty_pool(docx_path):
    """
    Extracts faculty names from the duty-pool Word document.
    """
    teachers = []
    if not os.path.exists(docx_path):
        return teachers
    
    doc = Document(docx_path)
    for table in doc.tables:
        for row in table.rows[1:]:
            cells = [c.text.strip() for c in row.cells]
            if len(cells) >= 2 and cells[1]:
                name = cells[1].replace('\n', ' ').strip()
                if name and name not in ['Dr. Sandeep Kad', 'Er. Ajay Sharma', 'Name']:
                    if name not in teachers:
                        teachers.append(name)
    return teachers


def generate_seating_plan(output_path, students_df, faculty_list, blueprint_rooms, exam_info):
    """
    Generates structured examination seating plan workbook in Excel.
    - Strictly 1 invigilator allocated per room.
    - Institutional header set to Amritsar Group of Colleges.
    - Dynamic interleaved allocation supporting 3rd, 5th, and 7th semesters.
    """
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # Remove default blank sheet

    font_title = Font(name="Arial", size=13, bold=True)
    font_bold = Font(name="Arial", size=10, bold=True)
    font_regular = Font(name="Arial", size=10)
    
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    right_align = Alignment(horizontal="right", vertical="center")
    
    thin_side = Side(border_style="thin", color="000000")
    box_border = Border(top=thin_side, left=thin_side, right=thin_side, bottom=thin_side)

    # Filter only eligible students
    eligible_df = students_df[students_df['Status'].astype(str).str.lower() == 'eligible'].copy()
    active_sems = sorted(list(eligible_df['Semester'].unique()))
    sem_pools = {s: eligible_df[eligible_df['Semester'] == s].to_dict('records') for s in active_sems}
    sem_ptrs = {s: 0 for s in active_sems}

    faculty_available = [f for f in faculty_list]

    college_name = exam_info.get("college_name", "Amritsar Group of Colleges")
    exam_name = exam_info.get("exam_name", "End Sem Examination")
    exam_date = exam_info.get("exam_date", "06/05/2026")
    exam_time = exam_info.get("exam_time", "10:30 am to 12:30 pm")
    sub_details = exam_info.get("sub_details", {})

    def get_semester_student(target_sem):
        if target_sem in sem_pools and sem_ptrs[target_sem] < len(sem_pools[target_sem]):
            st_obj = sem_pools[target_sem][sem_ptrs[target_sem]]
            sem_ptrs[target_sem] += 1
            return st_obj
        return None

    for r_item in blueprint_rooms:
        if all(sem_ptrs[s] >= len(sem_pools[s]) for s in active_sems):
            break

        room_no = r_item.get("room_no", "MB-102")
        ws = wb.create_sheet(title=f"Room-{room_no}")
        is_eclass = ('E-CLASS' in str(room_no).upper())

        # Set column widths
        for col in ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H', 'I', 'J', 'K', 'L']:
            ws.column_dimensions[col].width = 14
        ws.column_dimensions['A'].width = 8

        # --- 1. HEADER SECTION ---
        end_c = 10 if is_eclass else 12
        headers_data = [
            (1, "EXAMINATION SEATING PLAN", 13, True),
            (2, "Academic Examination Cell", 10, True),
            (3, college_name, 10, True),
            (4, exam_name, 11, True),
            (5, f"Room No:-  {room_no}", 12, True)
        ]

        for r_idx, text_val, f_size, is_bold in headers_data:
            ws.merge_cells(start_row=r_idx, start_column=2, end_row=r_idx, end_column=end_c)
            cell = ws.cell(r_idx, 2, value=text_val)
            cell.font = Font(name="Arial", size=f_size, bold=is_bold)
            cell.alignment = center_align

        ws.merge_cells("B6:E6")
        ws["B6"] = f"Time :- {exam_time}"
        ws["B6"].font = font_bold
        ws["B6"].alignment = left_align

        ws.merge_cells("H6:K6")
        ws["H6"] = f"Date :{exam_date} "
        ws["H6"].font = font_bold
        ws["H6"].alignment = right_align

        # --- 2. ROW BLOCKS & COLUMN HEADERS ---
        ws.merge_cells("B7:D7"); ws["B7"] = "Row 1"; ws["B7"].font = Font(name="Arial", size=14, bold=True); ws["B7"].alignment = center_align
        ws.merge_cells("E7:I7"); ws["E7"] = "Row 2"; ws["E7"].font = Font(name="Arial", size=14, bold=True); ws["E7"].alignment = center_align
        ws.merge_cells("J7:L7"); ws["J7"] = "Row 3"; ws["J7"].font = Font(name="Arial", size=14, bold=True); ws["J7"].alignment = center_align

        # Interleaved column distribution based on active semesters
        if len(active_sems) >= 3:
            sA, sB, sC = active_sems[0], active_sems[1], active_sems[2]
            col_mapping = [(2, sA), (4, sB), (5, sA), (7, sC), (9, sB), (10, sA), (12, sB)]
        elif len(active_sems) == 2:
            sA, sB = active_sems[0], active_sems[1]
            col_mapping = [(2, sA), (4, sB), (5, sA), (7, sB), (10, sA), (12, sB)]
        else:
            sA = active_sems[0]
            col_mapping = [(2, sA), (4, sA), (5, sA), (7, sA), (10, sA), (12, sA)]

        room_students = []
        col_sections = {c: set() for c in range(1, 13)}

        for b in range(1, 8):
            ws.cell(8 + b, 1, value=b).border = box_border
            ws.cell(8 + b, 1).alignment = center_align

        for col_idx, designated_sem in col_mapping:
            for bench in range(1, 8):
                row_idx = 8 + bench
                st_item = get_semester_student(designated_sem)
                if st_item:
                    ws.cell(row_idx, col_idx, value=st_item['Roll No']).alignment = center_align
                    col_sections[col_idx].add(st_item['Class'])
                    room_students.append(st_item)
                else:
                    ws.cell(row_idx, col_idx, value=None)
                ws.cell(row_idx, col_idx).border = box_border

        ws.cell(8, 1, value="S.No.").font = font_bold
        ws.cell(8, 1).border = box_border
        ws.cell(8, 1).alignment = center_align

        for col in range(2, 13):
            if col_sections[col]:
                names = " / ".join([s.replace('B.Tech. ', '') for s in col_sections[col]])
                ws.cell(8, col, value=f"B.Tech. {names}").font = Font(name="Arial", size=8, bold=True)
            else:
                ws.cell(8, col, value=None)
            ws.cell(8, col).border = box_border
            ws.cell(8, col).alignment = center_align

        # --- 3. SUMMARY TABLE ---
        sum_start = 17
        summary_headers = ['Sr. no', 'Class', 'Total', 'Subject Code', 'Subject name', 'Present', 'Absent', 'Not Eligible']
        for c_idx, h in enumerate(summary_headers, 1):
            c_cell = ws.cell(sum_start, c_idx, value=h)
            c_cell.font = font_bold
            c_cell.alignment = center_align
            c_cell.border = box_border

        sem_summary = {}
        for st_item in room_students:
            s_sem = st_item['Semester']
            if s_sem not in sem_summary:
                sem_summary[s_sem] = {'count': 0, 'sections': set()}
            sem_summary[s_sem]['count'] += 1
            if st_item.get('Section'):
                sem_summary[s_sem]['sections'].add(st_item['Section'])

        r_start = sum_start + 1
        sr_counter = 1
        for s_sem in sorted(sem_summary.keys()):
            info = sem_summary[s_sem]
            sec_text = " / ".join(sorted(info['sections']))
            class_display = f"B.Tech. {s_sem} Sem ({sec_text})" if sec_text else f"B.Tech. {s_sem} Sem"
            current_sub = sub_details.get(s_sem, {'name': '', 'code': ''})
            
            row_vals = [sr_counter, class_display, info['count'], current_sub['code'], current_sub['name'], info['count'], "", ""]
            for col_i, val in enumerate(row_vals, 1):
                c_cell = ws.cell(r_start, col_i, value=val)
                c_cell.font = font_regular
                c_cell.alignment = center_align
                c_cell.border = box_border
            r_start += 1
            sr_counter += 1

        # Total Row
        total_in_room = len(room_students)
        for c_idx in range(1, 9):
            c_cell = ws.cell(r_start, c_idx)
            c_cell.font = font_bold
            c_cell.alignment = center_align
            c_cell.border = box_border
            if c_idx == 1:
                c_cell.value = "Total"
            elif c_idx in [3, 6]:
                c_cell.value = total_in_room

        # --- 4. EXACTLY 1 INVIGILATOR PER ROOM ---
        assigned_teacher = faculty_available.pop(0) if faculty_available else "Faculty TBA"
        invig_row = r_start + 2
        invig_cell = ws.cell(invig_row, 1, value=f"Invigilator: {assigned_teacher}")
        invig_cell.font = Font(name="Arial", size=11, bold=True)
        invig_cell.alignment = left_align

    wb.save(output_path)
    return output_path