import streamlit as st
import pandas as pd
import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Font, Alignment, Border, Side
from docx import Document
import os
import math
import time
import io
import atexit
import shutil

# Page Configuration
st.set_page_config(page_title="Exam Seating Plan Generator", layout="wide")

OUTPUT_DIR = os.path.abspath('output')

def cleanup_output_files():
    if os.path.exists(OUTPUT_DIR):
        try:
            shutil.rmtree(OUTPUT_DIR)
        except Exception:
            pass

atexit.register(cleanup_output_files)

if 'initialized' not in st.session_state:
    cleanup_output_files()
    st.session_state.initialized = True

def parse_faculty_docx(file_source):
    faculty = []
    try:
        doc = Document(file_source)
        for table in doc.tables:
            for row in table.rows[1:]:
                cells = [c.text.strip() for c in row.cells]
                if len(cells) >= 2 and cells[1]:
                    name = cells[1].replace('\n', ' ').strip()
                    if name and name not in ['Dr. Sandeep Kad', 'Er. Ajay Sharma', 'Name']:
                        if name not in faculty:
                            faculty.append(name)
    except Exception as e:
        st.error(f"Error reading faculty file: {e}")
    return faculty

def extract_students_from_sheet(df, filename, sheet_name=""):
    roll_col, name_col, header_row = None, None, None
    for r in range(min(15, df.shape[0])):
        for c in range(df.shape[1]):
            val = str(df.iloc[r, c]).strip().lower()
            if ('roll' in val or 'rollno' in val) and roll_col is None:
                roll_col = c
                header_row = r
            if ('name' in val or 'student' in val) and 'father' not in val and name_col is None:
                name_col = c

    f_lower = str(filename).strip().lower()
    s_lower = str(sheet_name).strip().lower()

    # 1. Semester Detection (Priority to Filename)
    sem = "Unknown"
    if '3rd' in f_lower or 'sem 3' in f_lower or '3rdd' in f_lower:
        sem = "3rd"
    elif '5th' in f_lower or 'sem 5' in f_lower or '-5' in f_lower or '_5' in f_lower:
        sem = "5th"
    elif '7th' in f_lower or 'sem 7' in f_lower or '-7' in f_lower or '_7' in f_lower:
        sem = "7th"
    else:
        if '3rd' in s_lower or 'sem 3' in s_lower:
            sem = "3rd"
        elif '5th' in s_lower or 'sem 5' in s_lower:
            sem = "5th"
        elif '7th' in s_lower or 'sem 7' in s_lower:
            sem = "7th"

    # 2. Section Detection (Sheet name takes first priority)
    sec = ""
    if 'coe' in s_lower:
        sec = "COE"
    elif '3(1)' in s_lower or '3-1' in s_lower or 'cse 3(1)' in s_lower or 'cse3(1)' in s_lower:
        sec = "CSE-I"
    elif '3(2)' in s_lower or '3-2' in s_lower or 'cse 3(2)' in s_lower or 'cse3(2)' in s_lower:
        sec = "CSE-II"
    elif '3(3)' in s_lower or '3-3' in s_lower:
        sec = "CSE-III"
    elif 'coe' in f_lower:
        sec = "COE"
    elif any(x in f_lower or x in s_lower for x in ['cse1', 'cse-1', 'cse 1', 'cse_1']):
        sec = "CSE-I"
    elif any(x in f_lower or x in s_lower for x in ['cse2', 'cse-2', 'cse 2', 'cse_2']):
        sec = "CSE-II"
    elif any(x in f_lower or x in s_lower for x in ['cse3', 'cse-3', 'cse 3', 'cse_3']):
        sec = "CSE-III"
    elif any(x in f_lower or x in s_lower for x in ['cse4', 'cse-4', 'cse 4', 'cse_4']):
        sec = "CSE-IV"
    else:
        sec = "Regular"

    class_label = f"B.Tech. {sem} Sem ({sec})" if sec else f"B.Tech. {sem} Sem"

    # 3. Record Extraction
    records = []
    if roll_col is not None and header_row is not None:
        for r in range(header_row + 1, df.shape[0]):
            r_val = df.iloc[r, roll_col]
            n_val = df.iloc[r, name_col] if name_col is not None else ""
            if pd.notna(r_val):
                sval = str(r_val).strip()
                if any(ch.isdigit() for ch in sval) and len(sval) >= 4 and 'sem' not in sval.lower():
                    try:
                        if float(sval) == int(float(sval)):
                            sval = str(int(float(sval)))
                    except Exception:
                        pass
                    records.append({
                        'Roll No': sval,
                        'Name': str(n_val).strip() if pd.notna(n_val) else '',
                        'Semester': sem,
                        'Class': class_label,
                        'Section': sec,
                        'Status': 'Eligible'
                    })
    return pd.DataFrame(records)

def extract_students(file_source, filename):
    try:
        excel_file = pd.ExcelFile(file_source)
        all_sheet_records = []
        for sheet_name in excel_file.sheet_names:
            df = pd.read_excel(file_source, sheet_name=sheet_name, header=None)
            sheet_df = extract_students_from_sheet(df, filename, sheet_name)
            if not sheet_df.empty:
                all_sheet_records.append(sheet_df)
        if all_sheet_records:
            return pd.concat(all_sheet_records, ignore_index=True)
    except Exception:
        pass
    return pd.DataFrame()

def safe_write(ws, row, col, val):
    cell = ws.cell(row, col)
    if isinstance(cell, MergedCell):
        for rng in ws.merged_cells.ranges:
            if row >= rng.min_row and row <= rng.max_row and col >= rng.min_col and col <= rng.max_col:
                ws.cell(rng.min_row, rng.min_col).value = val
                return
    else:
        cell.value = val

st.title("📋 Automated Exam Seating Plan Generator")

if 'students_data' not in st.session_state:
    st.session_state.students_data = pd.DataFrame(columns=['Roll No', 'Name', 'Semester', 'Class', 'Section', 'Status'])

if 'faculty_list' not in st.session_state:
    st.session_state.faculty_list = []

# Sidebar Exam Details
st.sidebar.header("📝 Exam Details")
exam_name = st.sidebar.text_input("Exam Name", "End Sem Examination")
exam_date = st.sidebar.text_input("Exam Date", "06/05/2026")
exam_time = st.sidebar.text_input("Exam Time", "10:30 am to 12:30 pm")

# File Uploaders
st.subheader("📤 Upload Examination Configuration Files")
up1, up2, up3 = st.columns(3)

with up1:
    up_students = st.file_uploader("1. Student Lists (.xlsx)", type=["xlsx"], accept_multiple_files=True)
    if up_students:
        all_dfs = []
        for f in up_students:
            df_i = extract_students(f, f.name)
            if not df_i.empty:
                all_dfs.append(df_i)
        if all_dfs:
            st.session_state.students_data = pd.concat(all_dfs, ignore_index=True)

with up2:
    up_faculty = st.file_uploader("2. Faculty Duty Pool (.docx)", type=["docx"])
    if up_faculty:
        st.session_state.faculty_list = parse_faculty_docx(up_faculty)

with up3:
    up_template = st.file_uploader("3. Classroom Layout / Blueprint Sheet (.xlsx)", type=["xlsx"])

default_template_path = 'input_files/6 may 2026.xlsx'
if not up_template and os.path.exists(default_template_path):
    active_template_file = default_template_path
elif up_template:
    active_template_file = up_template
else:
    active_template_file = None

if st.session_state.students_data.empty:
    st.info("👆 Please upload student Excel sheets to continue.")
    st.stop()

# Detected Semesters
all_detected_sems = [s for s in sorted(list(st.session_state.students_data['Semester'].unique())) if s != "Unknown"]

st.sidebar.subheader("🎯 Active Exam Semesters")
selected_sems = st.sidebar.multiselect(
    "Active Semesters:",
    options=all_detected_sems,
    default=all_detected_sems
)

if not selected_sems:
    st.warning("Please select at least one semester!")
    st.stop()

# Subject Inputs
st.sidebar.subheader("📚 Subject Details")
sub_details = {}
default_subs = {
    "3rd": {"name": "Data Structures", "code": "AGCS – 21301"},
    "5th": {"name": "Operating System", "code": "AGCS – 21404"},
    "7th": {"name": "Big Data / DS", "code": "AGCS – 21603"}
}

for s in selected_sems:
    st.sidebar.markdown(f"**{s} Semester**")
    c1, c2 = st.sidebar.columns(2)
    def_info = default_subs.get(s, {"name": f"Subject {s}", "code": f"AGCS – {s}"})
    s_name = c1.text_input(f"Subject ({s})", value=def_info["name"], key=f"inp_sub_{s}")
    s_code = c2.text_input(f"Code ({s})", value=def_info["code"], key=f"inp_code_{s}")
    sub_details[s] = {'name': s_name, 'code': s_code}

filtered_students = st.session_state.students_data[st.session_state.students_data['Semester'].isin(selected_sems)].copy()

tab1, tab2, tab3 = st.tabs(["👥 Student Eligibility", "📊 Room Estimation & Generate", "👨‍🏫 Faculty Pool Preview"])

with tab1:
    st.subheader("Student Eligibility List")
    classes = ["All"] + sorted(list(filtered_students['Class'].unique()))
    sel_class = st.selectbox("Filter Class:", classes)
    
    view_df = filtered_students if sel_class == "All" else filtered_students[filtered_students['Class'] == sel_class]
    
    edited_df = st.data_editor(
        view_df,
        column_config={
            "Status": st.column_config.SelectboxColumn("Eligibility", options=["Eligible", "Not Eligible"], required=True)
        },
        disabled=["Roll No", "Name", "Semester", "Class", "Section"],
        use_container_width=True,
        key="editor"
    )
    
    # Session state update with modified status
    if not edited_df.empty:
        for idx, row in edited_df.iterrows():
            st.session_state.students_data.loc[st.session_state.students_data['Roll No'] == row['Roll No'], 'Status'] = row['Status']

eligible_df = st.session_state.students_data[
    (st.session_state.students_data['Status'] == 'Eligible') & 
    (st.session_state.students_data['Semester'].isin(selected_sems))
].copy()

not_eligible_df = st.session_state.students_data[
    (st.session_state.students_data['Status'] == 'Not Eligible') & 
    (st.session_state.students_data['Semester'].isin(selected_sems))
].copy()

with tab3:
    st.subheader("Available Teachers for Duty")
    if st.session_state.faculty_list:
        st.write(f"Total Eligible Teachers: **{len(st.session_state.faculty_list)}**")
        f_df = pd.DataFrame({"Sr. No": range(1, len(st.session_state.faculty_list)+1), "Teacher Name": st.session_state.faculty_list})
        st.dataframe(f_df, use_container_width=True, hide_index=True)
    else:
        st.warning("No faculty file loaded yet.")

with tab2:
    st.subheader("Room Estimation")
    metrics_cols = st.columns(len(selected_sems) + 1)
    for i, s in enumerate(selected_sems):
        count = len(eligible_df[eligible_df['Semester'] == s])
        metrics_cols[i].metric(f"{s} Sem Eligible", count)
    
    total_eligible = len(eligible_df)
    room_cap = 49 if len(selected_sems) >= 3 else 42
    estimated_rooms = math.ceil(total_eligible / room_cap) if room_cap else 1
    metrics_cols[-1].metric("Total Rooms Required", estimated_rooms)

    if st.button("🚀 Generate & Open Seating Plan", type="primary"):
        if active_template_file is None:
            st.error("Classroom Layout / Blueprint Sheet Excel file not found! Please upload above.")
            st.stop()

        faculty_available = [f for f in st.session_state.faculty_list]

        if isinstance(active_template_file, str):
            wb = openpyxl.load_workbook(active_template_file)
        else:
            wb = openpyxl.load_workbook(io.BytesIO(active_template_file.getvalue()))

        all_template_rooms = [sname for sname in wb.sheetnames]
        sorted_sems = sorted(selected_sems)
        sem_pools = {s: eligible_df[eligible_df['Semester'] == s].to_dict('records') for s in sorted_sems}
        sem_ptrs = {s: 0 for s in sorted_sems}

        thin = Side(border_style="thin", color="000000")
        box_border = Border(left=thin, right=thin, top=thin, bottom=thin)
        no_border = Border()

        used_rooms = []

        def get_next_student(preferred_sem=None, avoid_sem=None):
            if preferred_sem and sem_ptrs.get(preferred_sem, 0) < len(sem_pools.get(preferred_sem, [])):
                st_obj = sem_pools[preferred_sem][sem_ptrs[preferred_sem]]
                sem_ptrs[preferred_sem] += 1
                return st_obj
            for alt_sem in sorted_sems:
                if alt_sem != avoid_sem and sem_ptrs[alt_sem] < len(sem_pools[alt_sem]):
                    st_obj = sem_pools[alt_sem][sem_ptrs[alt_sem]]
                    sem_ptrs[alt_sem] += 1
                    return st_obj
            for alt_sem in sorted_sems:
                if sem_ptrs[alt_sem] < len(sem_pools[alt_sem]):
                    st_obj = sem_pools[alt_sem][sem_ptrs[alt_sem]]
                    sem_ptrs[alt_sem] += 1
                    return st_obj
            return None

        for room_order, r_name in enumerate(all_template_rooms):
            if all(sem_ptrs[s] >= len(sem_pools[s]) for s in sorted_sems):
                break

            used_rooms.append(r_name)
            ws = wb[r_name]
            is_eclass = ('E-CLASS' in r_name.upper())

            # Institutional Header
            h1 = "EXAMINATION SEATING PLAN"
            h2 = "Academic Examination Cell"
            h3 = "Amritsar Group of Colleges"
            h4 = exam_name
            h5 = f"Room No:-  {r_name}"

            for rng in [r for r in list(ws.merged_cells.ranges) if r.min_row <= 5]:
                ws.unmerge_cells(str(rng))

            for r in range(1, 6):
                for c in range(1, 14):
                    ws.cell(r, c).value = None

            end_c = 10 if is_eclass else 12
            headers_data = [(1, h1, 13, True), (2, h2, 10, True), (3, h3, 10, True), (4, h4, 11, True), (5, h5, 12, True)]
            for r_idx, text_val, f_size, is_bold in headers_data:
                ws.merge_cells(start_row=r_idx, start_column=2, end_row=r_idx, end_column=end_c)
                c_cell = ws.cell(r_idx, 2)
                c_cell.value = text_val
                c_cell.font = Font(name="Arial", size=f_size, bold=is_bold)
                c_cell.alignment = Alignment(horizontal="center", vertical="center")

            safe_write(ws, 6, 2, f"Time :- {exam_time}")
            safe_write(ws, 6, 8 if is_eclass else 10, f"Date :{exam_date} ")

            for rng in [r for r in list(ws.merged_cells.ranges) if r.min_row >= 7]:
                ws.unmerge_cells(str(rng))

            for r in range(7, 35):
                for c in range(1, 14):
                    ws.cell(r, c).value = None
                    ws.cell(r, c).border = no_border

            room_students = []

            if is_eclass:
                for i, c in enumerate([2, 4, 6, 8, 10], 1):
                    ws.cell(7, c).value = f"Row {i}"
                    ws.cell(7, c).font = Font(name="Arial", size=12, bold=True)
                    ws.cell(7, c).alignment = Alignment(horizontal="center", vertical="center")

                ws.cell(8, 1).value = "S.No."
                for b in range(1, 11):
                    ws.cell(8 + b, 1).value = b
                    ws.cell(8 + b, 1).border = box_border

                col_sections = {2: set(), 4: set(), 6: set(), 8: set(), 10: set()}
                cols_order = [2, 4, 6, 8, 10]

                last_col_sem = None
                for i_col, c_pos in enumerate(cols_order):
                    target_sem = sorted_sems[i_col % len(sorted_sems)]
                    for bench in range(1, 11):
                        row_idx = 8 + bench
                        st_item = get_next_student(preferred_sem=target_sem, avoid_sem=last_col_sem)
                        if st_item:
                            safe_write(ws, row_idx, c_pos, st_item['Roll No'])
                            col_sections[c_pos].add(st_item['Class'])
                            room_students.append(st_item)
                            last_col_sem = st_item['Semester']
                        else:
                            safe_write(ws, row_idx, c_pos, None)

                for col in cols_order:
                    if col_sections[col]:
                        names = " / ".join([s.replace('B.Tech. ', '') for s in col_sections[col]])
                        ws.cell(8, col).value = f"B.Tech. {names}"
                        ws.cell(8, col).font = Font(name="Arial", size=8, bold=True)

                for r in range(8, 19):
                    for c in [1, 2, 4, 6, 8, 10]:
                        ws.cell(r, c).border = box_border
                        ws.cell(r, c).alignment = Alignment(horizontal="center", vertical="center")

                sum_start = 20

            else:
                ws.merge_cells(start_row=7, start_column=2, end_row=7, end_column=4)
                ws.cell(7, 2).value = "Row 1"
                ws.cell(7, 2).font = Font(name="Arial", size=14, bold=True)
                ws.cell(7, 2).alignment = Alignment(horizontal="center", vertical="center")

                if len(sorted_sems) >= 3:
                    ws.merge_cells(start_row=7, start_column=5, end_row=7, end_column=9)
                else:
                    ws.merge_cells(start_row=7, start_column=5, end_row=7, end_column=7)

                ws.cell(7, 5).value = "Row 2"
                ws.cell(7, 5).font = Font(name="Arial", size=14, bold=True)
                ws.cell(7, 5).alignment = Alignment(horizontal="center", vertical="center")

                ws.merge_cells(start_row=7, start_column=10, end_row=7, end_column=12)
                ws.cell(7, 10).value = "Row 3"
                ws.cell(7, 10).font = Font(name="Arial", size=14, bold=True)
                ws.cell(7, 10).alignment = Alignment(horizontal="center", vertical="center")

                ws.column_dimensions['A'].width = 6
                ws.column_dimensions['B'].width = 14
                ws.column_dimensions['C'].width = 10
                ws.column_dimensions['D'].width = 14
                ws.column_dimensions['E'].width = 14
                ws.column_dimensions['F'].width = 10
                ws.column_dimensions['G'].width = 14
                ws.column_dimensions['H'].width = 10
                ws.column_dimensions['I'].width = 14
                ws.column_dimensions['J'].width = 14
                ws.column_dimensions['K'].width = 10
                ws.column_dimensions['L'].width = 14

                for b in range(1, 8):
                    ws.cell(8 + b, 1).value = b
                    ws.cell(8 + b, 1).border = box_border
                    ws.cell(8 + b, 1).alignment = Alignment(horizontal="center", vertical="center")

                col_sections = {c: set() for c in range(1, 13)}

                if len(sorted_sems) >= 3:
                    sem_A, sem_B, sem_C = sorted_sems[0], sorted_sems[1], sorted_sems[2]
                    col_mapping = [
                        (2, sem_A),
                        (4, sem_B),
                        (5, sem_A),
                        (7, sem_C),
                        (9, sem_B),
                        (10, sem_A),
                        (12, sem_B)
                    ]
                else:
                    sem_A = sorted_sems[0]
                    sem_B = sorted_sems[1] if len(sorted_sems) > 1 else sorted_sems[0]
                    col_mapping = [
                        (2, sem_A),
                        (4, sem_B),
                        (5, sem_A),
                        (7, sem_B),
                        (10, sem_A),
                        (12, sem_B)
                    ]

                prev_sem = None
                for col_idx, designated_sem in col_mapping:
                    for bench in range(1, 8):
                        row_idx = 8 + bench
                        st_item = get_next_student(preferred_sem=designated_sem, avoid_sem=prev_sem)
                        if st_item:
                            safe_write(ws, row_idx, col_idx, st_item['Roll No'])
                            col_sections[col_idx].add(st_item['Class'])
                            room_students.append(st_item)
                            prev_sem = st_item['Semester']
                        else:
                            safe_write(ws, row_idx, col_idx, None)

                for r in range(8, 16):
                    for c in range(1, 13):
                        ws.cell(r, c).border = box_border
                        ws.cell(r, c).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

                ws.cell(8, 1).value = "S.No."
                for col in range(2, 13):
                    if col_sections[col]:
                        names = " / ".join([s.replace('B.Tech. ', '') for s in col_sections[col]])
                        ws.cell(8, col).value = f"B.Tech. {names}"
                        ws.cell(8, col).font = Font(name="Arial", size=8, bold=True)
                    else:
                        ws.cell(8, col).value = None

                sum_start = 17

            # Summary Table
            ws.row_dimensions[sum_start].height = 28
            summary_headers = ['Sr. no', 'Class', 'Total', 'Subject Code', 'Subject name', 'Present', 'Absent', 'Not Eligible']
            for c_idx, h in enumerate(summary_headers, 1):
                c_cell = ws.cell(sum_start, c_idx)
                c_cell.value = h
                c_cell.font = Font(name="Arial", size=10, bold=True)
                c_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                c_cell.border = box_border

            # Class-wise accurate grouping
            class_counts = {}
            for st_item in room_students:
                cls_name = st_item['Class']
                s_sem = st_item['Semester']
                if cls_name not in class_counts:
                    class_counts[cls_name] = {'count': 0, 'sem': s_sem}
                class_counts[cls_name]['count'] += 1

            r_start = sum_start + 1
            sr_counter = 1
            total_not_eligible_room = 0

            for cls_name, info in class_counts.items():
                s_sem = info['sem']
                # Accurate Not Eligible count
                ne_count = len(not_eligible_df[not_eligible_df['Class'] == cls_name])
                total_not_eligible_room += ne_count
                
                ws.row_dimensions[r_start].height = 24
                current_sub = sub_details.get(s_sem, {'name': '', 'code': ''})
                
                row_data = [
                    sr_counter, 
                    cls_name, 
                    info['count'], 
                    current_sub['code'], 
                    current_sub['name'], 
                    info['count'], 
                    "", 
                    ne_count if ne_count > 0 else ""
                ]
                for col_i, val in enumerate(row_data, 1):
                    c_cell = ws.cell(r_start, col_i)
                    c_cell.value = val
                    c_cell.font = Font(name="Arial", size=10)
                    c_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                    c_cell.border = box_border

                r_start += 1
                sr_counter += 1

            # Total Row
            ws.row_dimensions[r_start].height = 24
            total_in_room = len(room_students)
            for c_idx in range(1, 9):
                c_cell = ws.cell(r_start, c_idx)
                c_cell.font = Font(name="Arial", size=10, bold=True)
                c_cell.alignment = Alignment(horizontal="center", vertical="center")
                c_cell.border = box_border
                if c_idx == 1:
                    c_cell.value = "Total"
                elif c_idx in [3, 6]:
                    c_cell.value = total_in_room
                elif c_idx == 8:
                    c_cell.value = total_not_eligible_room if total_not_eligible_room > 0 else ""
                else:
                    c_cell.value = None

            for r_below in range(r_start + 1, r_start + 10):
                for c_below in range(1, 13):
                    ws.cell(r_below, c_below).border = no_border
                    ws.cell(r_below, c_below).value = None

            # Strictly 1 Invigilator per room
            assigned_teacher = faculty_available.pop(0) if faculty_available else "Faculty TBA"
            invig_row = r_start + 2
            ws.row_dimensions[invig_row].height = 24
            invig_cell = ws.cell(invig_row, 1)
            invig_cell.value = f"Invigilator: {assigned_teacher}"
            invig_cell.font = Font(name="Arial", size=11, bold=True)
            invig_cell.alignment = Alignment(horizontal="left", vertical="center")
            invig_cell.border = no_border

        # Remove Unused Rooms
        sheets_in_wb = list(wb.sheetnames)
        for sname in sheets_in_wb:
            if sname not in used_rooms:
                wb.remove(wb[sname])

        # --- FORCE FIRST ROOM (START SE) OPEN HONA (Error-Free) ---
        if used_rooms:
            for ws_item in wb.worksheets:
                ws_item.sheet_view.tabSelected = False

            wb.active = 0
            wb.worksheets[0].sheet_view.tabSelected = True

        os.makedirs(OUTPUT_DIR, exist_ok=True)
        out_file = os.path.join(OUTPUT_DIR, 'Final_Seating_Plan.xlsx')
        
        excel_buffer = io.BytesIO()
        wb.save(excel_buffer)
        excel_buffer.seek(0)
        
        try:
            wb.save(out_file)
        except PermissionError:
            out_file = os.path.join(OUTPUT_DIR, f'Final_Seating_Plan_{int(time.time())}.xlsx')
            wb.save(out_file)

        st.success("Plan Generated Successfully!")
        
        st.download_button(
            label="📥 Download Final Seating Plan (Excel)",
            data=excel_buffer,
            file_name=f"Final_Seating_Plan_{exam_date.replace('/', '-')}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        
        try:
            os.startfile(out_file)
        except Exception:
            pass